from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from backend.services.dokument_scanner import sicherer_pfad


class DokumentLesefehler(Exception):
    pass


def _verstaendliche_fehlermeldung(datei: Path) -> str:
    bezeichnungen = {
        ".pdf": "PDF-Datei",
        ".docx": "Word-Dokument",
        ".xlsx": "Excel-Arbeitsmappe",
        ".txt": "Textdatei",
        ".pptx": "PowerPoint-Präsentation",
        ".ppsx": "PowerPoint-Bildschirmpräsentation",
    }
    bezeichnung = bezeichnungen.get(datei.suffix.lower(), "Dokument")
    return f"{bezeichnung} konnte nicht gelesen werden. Die Datei ist möglicherweise beschädigt oder nicht unterstützt."


def _pdf(pfad: Path) -> str:
    return "\n\n".join((seite.extract_text() or "") for seite in PdfReader(str(pfad)).pages)


def _docx(pfad: Path) -> str:
    doc = Document(str(pfad))
    teile = [p.text for p in doc.paragraphs if p.text]
    for tabelle in doc.tables:
        teile.extend("\t".join(zelle.text for zelle in zeile.cells) for zeile in tabelle.rows)
    return "\n".join(teile)


def _xlsx(pfad: Path) -> str:
    # data_only=False liest Formelausdrücke, ohne sie auszuführen.
    workbook = load_workbook(filename=str(pfad), read_only=True, data_only=False)
    try:
        teile = []
        for blatt in workbook.worksheets:
            teile.append(f"[Tabellenblatt: {blatt.title}]")
            for zeile in blatt.iter_rows(values_only=True):
                werte = ["" if wert is None else str(wert) for wert in zeile]
                if any(werte):
                    teile.append("\t".join(werte))
        return "\n".join(teile)
    finally:
        workbook.close()


def _praesentation(pfad: Path) -> str:
    praesentation = Presentation(str(pfad))
    teile = []
    for nummer, folie in enumerate(praesentation.slides, 1):
        texte = [shape.text for shape in folie.shapes if hasattr(shape, "text_frame") and shape.text]
        teile.append(f"[Folie {nummer}]" + (("\n" + "\n".join(texte)) if texte else ""))
    return "\n\n".join(teile)


def lese_dokument(datei: Path) -> str:
    try:
        pfad = sicherer_pfad(datei)
        endung = pfad.suffix.lower()
        if endung == ".txt":
            return pfad.read_text(encoding="utf-8-sig", errors="replace")
        if endung == ".pdf":
            return _pdf(pfad)
        if endung == ".docx":
            return _docx(pfad)
        if endung == ".xlsx":
            return _xlsx(pfad)
        if endung in {".pptx", ".ppsx"}:
            return _praesentation(pfad)
        raise DokumentLesefehler("Nicht unterstützter Dateityp.")
    except DokumentLesefehler:
        raise
    except Exception as exc:
        raise DokumentLesefehler(_verstaendliche_fehlermeldung(datei)) from exc
